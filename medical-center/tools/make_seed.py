import openpyxl, json, sys
p=sys.argv[1]
wb=openpyxl.load_workbook(p,data_only=True)
ws=wb["room"]
def num(v):
    try: return float(v)
    except: return 0
seed={"services":[],"doctors":[],"companies":[],"staff":[]}
seen=set()
for r in range(7,60):
    n=ws.cell(r,9).value
    if n and n.strip() and n.strip()!="عيادات الزهراء":
        seed["services"].append({"type":"clinic","name":n.strip(),"price":num(ws.cell(r,10).value),"tax":num(ws.cell(r,11).value)})
clinics={s["name"] for s in seed["services"]}
for r in range(7,48):
    n=ws.cell(r,2).value; c=ws.cell(r,3).value
    if n and c and (n.strip(),c.strip()) not in seen and n.strip()!="عيادات الزهراء":
        seen.add((n.strip(),c.strip()))
        seed["doctors"].append({"name":n.strip(),"clinic":c.strip(),"percent":num(ws.cell(r,5).value),"fixed":num(ws.cell(r,6).value),"opPercent":num(ws.cell(r,7).value)})
for r in range(7,30):
    n=ws.cell(r,14).value
    if n: seed["companies"].append({"name":str(n).strip(),"fee":num(ws.cell(r,16).value)})
for r in range(7,40):
    n=ws.cell(r,19).value
    if n and str(n).strip(): seed["staff"].append({"name":str(n).strip()})
exp=[str(ws.cell(r,30).value).strip() for r in range(7,60) if ws.cell(r,30).value]
seed["expenseCats"]=list(dict.fromkeys(exp))
for d in seed["doctors"]:
    if d["clinic"] not in clinics:
        seed["services"].append({"type":"clinic","name":d["clinic"],"price":0,"tax":0}); clinics.add(d["clinic"])
# ---- كتالوج الخدمات من الإكسيل: قائمة الإيرادات (room!Y) + أسعار الخدمات المتكررة في شيت CO ----
def add(t,n,price=0,tax=0):
    if not any(x["type"]==t and x["name"]==n for x in seed["services"]): seed["services"].append({"type":t,"name":n,"price":price,"tax":tax})
# أشعة
for n,pr in [("أشعة أسنان (فيلم)",40),("أشعة أسنان (2 فيلم)",80),("أشعة أسنان (4 أفلام)",160),("أشعة عادية",0),("موجات صوتية",250)]: add("radiology",n,pr)
# تحاليل
for n in ["تحاليل","تحاليل شركات","معمل لاب تولاب"]: add("lab",n)
# عمليات
for n,pr in [("عمليات جراحية",0),("فتح خراج",400),("استئصال كيس شمعي",450),("استئصال بواسير وشرخ شرجي",0),("خياطة جرح تحت مخدر موضعي",77),("تركيب جبيرة تحت الركبة",110),("جبس تحت الركبة",0),("طوارئ جراحة",0),("طوارئ عظام",0)]: add("operation",n,pr)
# خدمات أخرى (من قائمة الإيرادات وشيت الشركات)
for n,pr in [("كشف",0),("استشارة",0),("غيار",0),("خياطة",0),("نسبة شركات",0),("خلع أسنان عادي (ضرس)",200),("خلع جراحي",350),("حشو أسنان",0),("حشو كمبوزيت (ضرس)",400),("حشو ليزر (ضرس)",400),("حشو عصب (ضرس)",750),("إزالة رواسب",300),("طربوش بورسلين",1000),("دعامة فايبر جلاس",600),("علاج طبيعي (جلسة)",50),("جلسات علاج طبيعي (6)",300),("جلسات علاج طبيعي (12)",600),("حجامة",200)]: add("service",n,pr)
json.dump(seed,open("tools/seed.json","w"),ensure_ascii=False,indent=1)
print({k:len(v) for k,v in seed.items()}, "missing clinics for doctors:", {d["clinic"] for d in seed["doctors"]}-clinics)
